"""장부 요약 → ledger_monthly (집 PC 일꾼 전용). 원천 두 갈래.

사장님이 매달 정리하는 장부는 구글 시트 '베어글스_장부'다. 이걸 읽는 길이
둘 있고, **되는 쪽을 알아서 쓴다**:

  ① 장부관리 폴더의 CSV  (로그인 필요 없음 · 기본)
     사장님이 시트에서 [파일 → 다운로드 → 쉼표로 구분된 값(.csv)] 해서
     장부관리 폴더에 넣어두면 그걸 읽는다. TOS 엑셀을 올리는 그 폴더라
     새 루틴이 아니다. 파일 이름에 '장부'만 들어가면 되고, 여러 개면 가장
     최근 것을 쓴다. 확정/예상 판단은 그 파일의 수정 시각으로 한다.

  ② 구글 시트 직접 읽기 (Drive API · OAuth token.json 필요)
     자동이라 더 좋다. 2026-09-09~13 은 OAuth 앱(beargels-sns)이 조직 전용으로
     잠겨 403 org_internal 로 막혀 있었고, 2026-09-13 사장님이 콘솔에서 풀어
     myeonggu96@gmail.com 으로 로그인한 뒤부터 돈다. 요약은 CSV 내보내기로,
     '매출장부' 시트는 xlsx 내보내기로 읽는다(CSV 는 첫 시트만 주니까).

그래서 순서는 **①을 먼저 보고, 없으면 ②**다. ①이 있으면 로그인이 영영
막혀 있어도 장부가 최신으로 유지된다.

일꾼의 하루 1회 매출 반영(maybe_pos_import)과 웹 [장부 지금 반영] 버튼에서
sync() 를 부른다. 실패해도 포스 장부 반영은 막지 않는다.
"""

from __future__ import annotations

import glob
import io
import logging
import os
import pathlib
import sys
from datetime import datetime, timezone

ROOT = pathlib.Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from database import ledger_store  # noqa: E402

logger = logging.getLogger(__name__)

# 시트 ID — 드라이브에서 '베어글스_장부' (사장님 계정). 바뀌면 .env 로.
DEFAULT_SHEET_ID = "1kPh9GkWP4g9qFLDlPKs0LbASx3Aj3dHkq9i8g2EHRe4"

# 장부관리 폴더에서 이 말이 이름에 든 .csv 를 장부로 본다.
# (구글 시트가 내려주는 기본 이름이 '베어글스_장부 - 요약.csv')
CSV_NAME_HINT = os.getenv("LEDGER_CSV_HINT", "장부")


def sheet_id() -> str:
    return os.getenv("LEDGER_SHEET_ID", DEFAULT_SHEET_ID)


def _token_file() -> str:
    p = os.getenv("GOOGLE_OAUTH_TOKEN_FILE", "token.json")
    return p if os.path.isabs(p) else str(ROOT / p)


# ---------------------------------------------------------------------------
# ① 장부관리 폴더의 CSV — 로그인 없이
# ---------------------------------------------------------------------------

def find_local_csv(base=None):
    """장부관리 폴더(하위 포함)에서 가장 최근 장부 CSV 경로. 없으면 None."""
    from worker import pos_import
    base = base or pos_import.ledger_dir()
    if not os.path.isdir(base):
        return None
    hits = [p for p in glob.glob(os.path.join(base, "**", "*.csv"), recursive=True)
            if CSV_NAME_HINT in os.path.basename(p)
            and not os.path.basename(p).startswith("~$")
            # 구글이 내려주는 다른 시트('베어글스_장부 - 매출장부.csv' 등)도
            # 이름에 '장부'가 든다 — 요약이 아닌 시트는 여기서 뺀다.
            and not any(x in os.path.basename(p) for x in OTHER_SHEETS)]
    return max(hits, key=os.path.getmtime) if hits else None


# 장부 파일의 다른 시트 이름 — 요약 CSV 를 고를 때 걸러내고, 매출장부는 따로 읽는다
OTHER_SHEETS = ("매출장부", "매입장부", "기타매입", "시간대별")


def find_local_sales_csv(base=None):
    """장부관리 폴더의 '매출장부' 시트 CSV(가장 최근). 없으면 None."""
    from worker import pos_import
    base = base or pos_import.ledger_dir()
    if not os.path.isdir(base):
        return None
    hits = [p for p in glob.glob(os.path.join(base, "**", "*.csv"), recursive=True)
            if "매출장부" in os.path.basename(p) and not os.path.basename(p).startswith("~$")]
    return max(hits, key=os.path.getmtime) if hits else None


def read_local_csv(path):
    """(csv_text, modified_at) — 수정 시각이 확정/예상을 가른다."""
    # 구글이 내려주는 CSV 는 UTF-8(BOM). 엑셀에서 다시 저장하면 cp949 일 수 있다.
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = io.open(path, encoding=enc).read()
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"장부 CSV 를 읽지 못했어요(글자 인코딩): {os.path.basename(path)}")
    mod = datetime.fromtimestamp(os.path.getmtime(path)).astimezone()
    return text, mod


# ---------------------------------------------------------------------------
# ② 구글 시트 직접 (OAuth)
# ---------------------------------------------------------------------------

def fetch_summary_csv():
    """(csv_text, modified_at: datetime|None). Drive API 두 번 왕복.

    files.export 는 **첫 번째 시트**(= 요약)만 내보내므로 요약시트가 맨 앞에
    있어야 한다(지금 그렇다).
    """
    from googleapiclient.discovery import build
    from sns_automation.drive_monitor import load_oauth_credentials

    creds = load_oauth_credentials(_token_file())
    svc = build("drive", "v3", credentials=creds, cache_discovery=False)
    fid = sheet_id()
    meta = svc.files().get(fileId=fid, fields="modifiedTime,name").execute()
    raw = svc.files().export(fileId=fid, mimeType="text/csv").execute()
    text = raw.decode("utf-8-sig") if isinstance(raw, bytes) else str(raw)
    mod = None
    if meta.get("modifiedTime"):
        mod = datetime.fromisoformat(meta["modifiedTime"].replace("Z", "+00:00"))
    return text, mod


def fetch_sales_rows():
    """구글 시트의 '매출장부' 시트 → 셀 값 행 목록. (xlsx 로 내보내 읽는다 —
    CSV 내보내기는 첫 시트만 주기 때문.)"""
    from googleapiclient.discovery import build
    from sns_automation.drive_monitor import load_oauth_credentials
    import openpyxl

    creds = load_oauth_credentials(_token_file())
    svc = build("drive", "v3", credentials=creds, cache_discovery=False)
    raw = svc.files().export(
        fileId=sheet_id(),
        mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ).execute()
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    if "매출장부" not in wb.sheetnames:
        return []
    return [list(r) for r in wb["매출장부"].iter_rows(values_only=True)]


def load_sales():
    """({ym: 매출장부 항목}, source) — 폴더 CSV 우선, 없으면 구글 시트.
    실패는 여기서 삼킨다: 매출장부 항목은 덤이고, 요약 반영을 막으면 안 된다."""
    import csv
    try:
        path = find_local_sales_csv()
        if path:
            text, _ = read_local_csv(path)
            return ledger_store.parse_sales_csv(text), f"폴더 CSV({os.path.basename(path)})"
        rows = fetch_sales_rows()
        return (ledger_store.parse_sales_rows(rows) if rows else {}), "구글 시트"
    except Exception as e:  # noqa: BLE001
        logger.warning("매출장부 시트 읽기 실패(요약은 계속): %s", e)
        return {}, f"실패({str(e)[:40]})"


# ---------------------------------------------------------------------------
# 반영
# ---------------------------------------------------------------------------

def load_summary():
    """(text, modified_at, source) — 폴더 CSV 우선, 없으면 구글 시트."""
    path = find_local_csv()
    if path:
        text, mod = read_local_csv(path)
        return text, mod, f"폴더 CSV({os.path.basename(path)})"
    text, mod = fetch_summary_csv()
    return text, mod, "구글 시트"


def sync(text=None, modified_at=None) -> dict:
    """장부 → ledger_monthly + 목표(menu_settings). 요약 dict 반환.

    text 를 주면(테스트·수동 시드) 아무것도 읽지 않는다.
    """
    source = "직접 넘김"
    if text is None:
        text, modified_at, source = load_summary()
    rows, targets = ledger_store.parse_summary_csv(text, modified_at)
    if text is not None and source != "직접 넘김":
        sales, sales_src = load_sales()          # 매출장부 시트 항목(덤)
        ledger_store.merge_sales(rows, sales)
        if sales:
            source += f" + 매출장부[{sales_src}]"
    n = ledger_store.upsert_ledger(rows)
    if targets:
        ledger_store.set_ledger_targets(targets)
    est = [r["ym"] for r in rows if r.get("status") == "estimate"]
    note = (f"{n}개월 반영 [{source}]"
            + (f" (예상치 {', '.join(est)})" if est else "")
            + (f", 장부 수정 {modified_at.astimezone(timezone.utc).date()}"
               if modified_at else ""))
    logger.info("장부 반영: %s", note)
    return {"ok": True, "months": n, "estimate": est, "targets": bool(targets),
            "source": source, "note": note}
